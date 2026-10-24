import csv
import io
import os
import re
from typing import List, Tuple
import openpyxl
import pypdf
from docx import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from app.schemas.document import ExtractionResult, SectionContent, TableData


def sanitize_filename(filename: str) -> str:
    """Sanitizes user filename to prevent path traversal and shell issues."""
    clean = os.path.basename(filename)
    clean = re.sub(r"[^\w\.\-\s]", "_", clean)
    clean = re.sub(r"\s+", " ", clean).strip()
    return clean or "unnamed_document"


def validate_file_magic_bytes(file_bytes: bytes, file_type: str) -> Tuple[bool, str]:
    """Validates magic bytes of uploaded content to ensure type integrity."""
    if len(file_bytes) == 0:
        return False, "File is empty (0 bytes)."

    if file_type == "pdf":
        if not file_bytes.startswith(b"%PDF-"):
            return False, "Invalid PDF: File does not start with valid PDF signature (%PDF-)."
    elif file_type in ("docx", "xlsx"):
        # OOXML documents are standard ZIP archives starting with PK\x03\x04
        if not file_bytes.startswith(b"PK\x03\x04"):
            return False, f"Invalid {file_type.upper()}: File is not a valid OpenXML archive."
    elif file_type in ("csv", "txt", "md"):
        # Plain text files must not contain binary null bytes in initial sample
        sample = file_bytes[:1024]
        if b"\x00" in sample:
            return False, f"Invalid {file_type.upper()}: File appears to be binary, not valid text."

    return True, ""


class DocumentExtractor:
    """Core document parsing engine for DevMatrix AI project documents."""

    @classmethod
    def extract(cls, file_path: str, original_filename: str, file_type: str, document_id: str) -> ExtractionResult:
        file_type = file_type.lower()
        warnings: List[str] = []
        sections: List[SectionContent] = []
        total_pages = 0
        total_sheets = 0
        status = "completed"

        try:
            if file_type == "pdf":
                sections, total_pages, warnings = cls._extract_pdf(file_path)
            elif file_type == "docx":
                sections, warnings = cls._extract_docx(file_path)
                total_pages = len(sections)
            elif file_type == "xlsx":
                sections, total_sheets, warnings = cls._extract_xlsx(file_path)
            elif file_type == "csv":
                sections, warnings = cls._extract_csv(file_path)
            elif file_type in ("txt", "md"):
                sections, warnings = cls._extract_text(file_path, file_type)
            else:
                return ExtractionResult(
                    document_id=document_id,
                    filename=original_filename,
                    file_type=file_type,
                    status="failed",
                    has_text=False,
                    warnings=[f"Unsupported file format: {file_type}"],
                    sections=[],
                    raw_preview="",
                )
        except Exception as exc:
            return ExtractionResult(
                document_id=document_id,
                filename=original_filename,
                file_type=file_type,
                status="failed",
                has_text=False,
                warnings=[f"Extraction failed: {str(exc)}"],
                sections=[],
                raw_preview="",
            )

        # Count total words across all legitimate text sections and table cells
        total_words = 0
        preview_snippets: List[str] = []

        for sec in sections:
            if sec.content:
                stripped = sec.content.strip()
                if stripped:
                    words = len(stripped.split())
                    total_words += words
                    if len(preview_snippets) < 3:
                        preview_snippets.append(stripped[:200])
            if sec.table:
                for row in sec.table.rows:
                    for cell in row:
                        if cell is not None and str(cell).strip():
                            total_words += len(str(cell).strip().split())

        has_text = total_words > 0
        if not has_text:
            status = "warning"
            warnings.append(
                "No extractable text found in file. The document may be empty, image-only, or scanned. "
                "(Note: OCR is not supported in this version)."
            )
        elif warnings:
            status = "warning"

        raw_preview = "\n\n".join(preview_snippets)[:500] if preview_snippets else ""

        return ExtractionResult(
            document_id=document_id,
            filename=original_filename,
            file_type=file_type,
            status=status,
            has_text=has_text,
            total_pages=total_pages,
            total_sheets=total_sheets,
            total_words=total_words,
            warnings=warnings,
            sections=sections,
            raw_preview=raw_preview,
        )

    @classmethod
    def _extract_pdf(cls, file_path: str) -> Tuple[List[SectionContent], int, List[str]]:
        sections: List[SectionContent] = []
        warnings: List[str] = []

        with open(file_path, "rb") as f:
            reader = pypdf.PdfReader(f)
            total_pages = len(reader.pages)

            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    warnings.append("PDF is password protected and could not be fully decrypted.")

            for idx, page in enumerate(reader.pages, start=1):
                try:
                    page_text = page.extract_text() or ""
                    clean_text = page_text.strip()
                    sections.append(
                        SectionContent(
                            type="page",
                            index=idx,
                            title=f"Page {idx}",
                            content=clean_text,
                        )
                    )
                except Exception as e:
                    warnings.append(f"Could not extract text from Page {idx}: {str(e)}")
                    sections.append(
                        SectionContent(
                            type="page",
                            index=idx,
                            title=f"Page {idx}",
                            content="",
                        )
                    )

        return sections, total_pages, warnings

    @classmethod
    def _extract_docx(cls, file_path: str) -> Tuple[List[SectionContent], List[str]]:
        sections: List[SectionContent] = []
        warnings: List[str] = []
        doc = DocxDocument(file_path)

        section_idx = 1
        for child in doc.element.body:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "p":
                para = DocxParagraph(child, doc)
                text = para.text.strip()
                if text:
                    sections.append(
                        SectionContent(
                            type="text",
                            index=section_idx,
                            title=f"Section {section_idx}",
                            content=text,
                        )
                    )
                    section_idx += 1
            elif tag == "tbl":
                table = DocxTable(child, doc)
                rows_data: List[List[str]] = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    rows_data.append(row_cells)

                if rows_data:
                    headers = rows_data[0]
                    content_rows = rows_data[1:] if len(rows_data) > 1 else []
                    sections.append(
                        SectionContent(
                            type="table",
                            index=section_idx,
                            title=f"Table {section_idx}",
                            table=TableData(headers=headers, rows=content_rows),
                        )
                    )
                    section_idx += 1

        return sections, warnings

    @classmethod
    def _extract_xlsx(cls, file_path: str) -> Tuple[List[SectionContent], int, List[str]]:
        sections: List[SectionContent] = []
        warnings: List[str] = []

        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        sheet_names = wb.sheetnames
        total_sheets = len(sheet_names)

        for idx, name in enumerate(sheet_names, start=1):
            sheet = wb[name]
            raw_rows: List[List[Any]] = []

            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None and str(cell).strip() != "" for cell in row):
                    cleaned_row = [str(c) if c is not None else "" for c in row]
                    raw_rows.append(cleaned_row)

            if raw_rows:
                headers = [str(h) for h in raw_rows[0]]
                body_rows = raw_rows[1:]
                sections.append(
                    SectionContent(
                        type="sheet",
                        index=idx,
                        title=f"Sheet: {name}",
                        table=TableData(headers=headers, rows=body_rows),
                    )
                )
            else:
                sections.append(
                    SectionContent(
                        type="sheet",
                        index=idx,
                        title=f"Sheet: {name}",
                        content="",
                    )
                )

        wb.close()
        return sections, total_sheets, warnings

    @classmethod
    def _extract_csv(cls, file_path: str) -> Tuple[List[SectionContent], List[str]]:
        sections: List[SectionContent] = []
        warnings: List[str] = []

        content_text = ""
        for encoding in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    content_text = f.read()
                break
            except UnicodeDecodeError:
                continue

        reader = csv.reader(io.StringIO(content_text))
        rows = [row for row in reader if any(cell.strip() for cell in row)]

        if rows:
            headers = rows[0]
            body_rows = rows[1:] if len(rows) > 1 else []
            sections.append(
                SectionContent(
                    type="table",
                    index=1,
                    title="CSV Data",
                    table=TableData(headers=headers, rows=body_rows),
                )
            )
        else:
            sections.append(
                SectionContent(
                    type="table",
                    index=1,
                    title="CSV Data",
                    content="",
                )
            )

        return sections, warnings

    @classmethod
    def _extract_text(cls, file_path: str, file_type: str) -> Tuple[List[SectionContent], List[str]]:
        sections: List[SectionContent] = []
        warnings: List[str] = []

        content = ""
        for encoding in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    content = f.read()
                break
            except UnicodeDecodeError:
                continue

        clean_text = content.strip()
        sections.append(
            SectionContent(
                type="text",
                index=1,
                title=f"{file_type.upper()} Document",
                content=clean_text,
            )
        )

        return sections, warnings
